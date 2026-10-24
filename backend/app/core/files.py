"""PDF / DOCX bytes -> plain text."""
import io


def read_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    elif name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Use PDF or DOCX.")
    text = text.strip()
    if len(text) < 50:
        raise ValueError("No readable text found (scanned PDF?). Try the DOCX version.")
    return text
